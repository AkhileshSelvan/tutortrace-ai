from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
import subprocess
from zipfile import ZipFile

import pymupdf
import pytest
from pptx import Presentation
from pptx.util import Inches

from scripts.create_demo_assets import make_pdf, make_pptx
from tutortrace.errors import DocumentLimitError, InvalidDocumentError, PasswordProtectedError, UnsupportedFileError
from tutortrace.ingestion import pdf, pptx, service
from tutortrace.ingestion.service import ingest
from tutortrace.models import digest
from tutortrace.persistence import Store
from tutortrace.retrieval.search import search


def test_pdf_page_citations_keep_blank_pages(store, pdf_bytes):
    result = ingest(store, "physics.pdf", pdf_bytes)
    assert result.document.unit_count == 3
    assert result.document.chunk_count == 2
    evidence = search(store, "gravity")[0]
    assert evidence.location == 3
    assert evidence.citation == "physics.pdf, page 3"
    assert evidence.document_id == digest(pdf_bytes)
    with store.connection() as db:
        assert db.execute("SELECT text FROM units WHERE location=2").fetchone()[0] == ""


def test_pptx_slide_citations_keep_blank_slides(store, pptx_bytes):
    result = ingest(store, "biology.pptx", pptx_bytes)
    assert result.document.unit_count == 3
    assert result.document.chunk_count == 2
    evidence = search(store, "chlorophyll")[0]
    assert evidence.location == 3
    assert evidence.citation == "biology.pptx, slide 3"
    assert evidence.source_type == "pptx"


@pytest.mark.parametrize("fixture,filename", [("pdf_bytes", "source.pdf"), ("pptx_bytes", "source.pptx")])
def test_original_preservation(store, request, fixture, filename):
    data = request.getfixturevalue(fixture)
    result = ingest(store, filename, data)
    assert store.original(result.document.id) == data
    assert result.document.byte_size == len(data)


def test_duplicate_bytes_renamed_and_restart(store, pdf_bytes):
    first = ingest(store, "first.pdf", pdf_bytes)
    reopened = Store(store.path)
    second = ingest(reopened, "renamed.PDF", pdf_bytes)
    assert second.duplicate
    assert second.document == first.document
    assert len(reopened.list_documents()) == 1
    assert len(search(reopened, "energy")) == 2
    assert [row["status"] for row in reopened.recent_attempts()] == ["duplicate", "indexed"]
    assert reopened.recent_attempts()[0]["filename"] == "renamed.PDF"


def test_same_filename_different_bytes_are_distinct(store, pdf_bytes):
    first = ingest(store, "lesson.pdf", pdf_bytes)
    second = ingest(store, "lesson.pdf", make_pdf(("Different lesson on algebra.",)))
    assert not second.duplicate
    assert first.document.id != second.document.id
    assert len(store.list_documents()) == 2


def test_concurrent_duplicate_ingestion_is_atomic(store, pdf_bytes):
    def upload(number):
        return ingest(Store(store.path), f"source-{number}.pdf", pdf_bytes)
    with ThreadPoolExecutor(max_workers=4) as executor:
        results = list(executor.map(upload, range(4)))
    assert sum(not result.duplicate for result in results) == 1
    assert len(store.list_documents()) == 1
    assert len(store.recent_attempts()) == 4
    assert len(search(store, "gravity")) == 1


@pytest.mark.parametrize("filename,data,error", [
    ("lesson.txt", b"energy", UnsupportedFileError),
    ("slides.ppt", b"legacy", UnsupportedFileError),
    ("broken.pdf", b"not a PDF", InvalidDocumentError),
    ("broken.pdf", b"%PDF-1.7\ngarbage", InvalidDocumentError),
    ("broken.pptx", b"not a ZIP", InvalidDocumentError),
    ("empty.pdf", b"", InvalidDocumentError),
    ("empty.pptx", b"", InvalidDocumentError),
])
def test_invalid_files_persist_failure_without_index(store, filename, data, error):
    with pytest.raises(error):
        ingest(store, filename, data)
    assert store.list_documents() == []
    assert search(store, "energy") == []
    attempt = store.recent_attempts()[0]
    assert attempt["status"] == "failed"
    assert attempt["error_code"] == error.code
    assert attempt["document_id"] is None


def test_non_presentation_zip_is_rejected(store):
    stream = BytesIO()
    with ZipFile(stream, "w") as archive:
        archive.writestr("hello.txt", "Not a presentation")
    with pytest.raises(InvalidDocumentError):
        ingest(store, "pretend.pptx", stream.getvalue())


@pytest.mark.parametrize("extension,builder", [("pdf", make_pdf), ("pptx", make_pptx)])
def test_empty_documents_preserved_but_not_searchable(store, extension, builder):
    data = builder(("", ""))
    result = ingest(store, f"blank.{extension}", data)
    assert result.document.status == "empty"
    assert result.document.unit_count == 2
    assert result.document.chunk_count == 0
    assert store.original(result.document.id) == data
    assert search(store, "energy") == []
    assert any("No extractable text" in warning for warning in result.document.warnings)
    assert ingest(store, f"blank.{extension}", data).duplicate


def test_zero_slide_presentation(store):
    result = ingest(store, "zero.pptx", make_pptx(()))
    assert result.document.status == "empty"
    assert result.document.unit_count == 0


def test_encrypted_pdf_rejected(store):
    with pymupdf.open(stream=make_pdf(), filetype="pdf") as document:
        encrypted = document.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="test-only", owner_pw="test-owner")
    with pytest.raises(PasswordProtectedError):
        ingest(store, "locked.pdf", encrypted)
    assert store.recent_attempts()[0]["error_code"] == "password_protected"


def test_upload_limit(store, pdf_bytes, monkeypatch):
    monkeypatch.setattr(service, "MAX_UPLOAD_BYTES", 10)
    with pytest.raises(DocumentLimitError):
        ingest(store, "oversize.pdf", pdf_bytes)


@pytest.mark.parametrize("module,fixture,extension", [(pdf, "pdf_bytes", "pdf"), (pptx, "pptx_bytes", "pptx")])
@pytest.mark.parametrize("setting,value", [("MAX_UNITS", 1), ("MAX_TEXT_CHARS", 10)])
def test_extraction_limits(store, request, monkeypatch, module, fixture, extension, setting, value):
    monkeypatch.setattr(module, setting, value)
    with pytest.raises(DocumentLimitError):
        ingest(store, f"oversize.{extension}", request.getfixturevalue(fixture))
    assert store.list_documents() == []
    assert store.recent_attempts()[0]["error_code"] == "document_limit"


@pytest.mark.parametrize("setting", ["MAX_ARCHIVE_ENTRIES", "MAX_EXPANDED_BYTES"])
def test_pptx_archive_limits(store, pptx_bytes, monkeypatch, setting):
    monkeypatch.setattr(pptx, setting, 1)
    with pytest.raises(DocumentLimitError):
        ingest(store, "oversize.pptx", pptx_bytes)


def test_pdf_worker_timeout_records_failure(store, pdf_bytes, monkeypatch):
    def timeout(*args, **kwargs):
        raise subprocess.TimeoutExpired(args[0], 60)
    monkeypatch.setattr(pdf.subprocess, "run", timeout)
    with pytest.raises(DocumentLimitError, match="60 seconds"):
        ingest(store, "slow.pdf", pdf_bytes)
    assert store.list_documents() == []
    assert store.recent_attempts()[0]["error_code"] == "document_limit"


def test_pdf_worker_crash_records_failure(store, pdf_bytes, monkeypatch):
    monkeypatch.setattr(pdf.subprocess, "run", lambda *args, **kwargs: subprocess.CompletedProcess(args[0], 1, b"", b"native crash"))
    with pytest.raises(InvalidDocumentError, match="worker failed"):
        ingest(store, "crash.pdf", pdf_bytes)
    assert store.list_documents() == []
    assert store.recent_attempts()[0]["status"] == "failed"


def test_pdf_preview_uses_exact_page_and_rejects_missing_page(pdf_bytes):
    first = pdf.render_pdf_page(pdf_bytes, 1)
    third = pdf.render_pdf_page(pdf_bytes, 3)
    assert first.startswith(b"\x89PNG\r\n\x1a\n")
    assert first != third
    with pytest.raises(InvalidDocumentError, match="does not exist"):
        pdf.render_pdf_page(pdf_bytes, 4)


def test_mislabeled_formats_are_rejected(store, pdf_bytes, pptx_bytes):
    with pytest.raises(InvalidDocumentError):
        ingest(store, "pretend.pdf", pptx_bytes)
    with pytest.raises(InvalidDocumentError):
        ingest(store, "pretend.pptx", pdf_bytes)
    ingest(store, "real.pdf", pdf_bytes)
    with pytest.raises(InvalidDocumentError):
        ingest(store, "renamed.pptx", pdf_bytes)


def test_filename_is_metadata_not_a_path(store, pdf_bytes):
    result = ingest(store, "../../private/lesson.pdf", pdf_bytes)
    assert result.document.filename == "lesson.pdf"
    assert not (store.path.parent / "lesson.pdf").exists()


def test_long_filename_preserves_valid_extension(store, pdf_bytes):
    result = ingest(store, "x" * 245 + ".pdf", pdf_bytes)
    assert len(result.document.filename) == 240
    assert result.document.filename.endswith(".pdf")
    assert result.document.status == "indexed"


def test_pptx_group_and_table_text(store):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(0, 0, Inches(3), Inches(1)).text_frame.text = "Grouped acceleration"
    table = slide.shapes.add_table(2, 2, 0, Inches(2), Inches(6), Inches(2)).table
    table.cell(0, 0).text = "Force"
    table.cell(0, 1).text = "newtons"
    table.cell(1, 0).text = "Mass"
    table.cell(1, 1).text = "kilograms"
    stream = BytesIO()
    presentation.save(stream)
    ingest(store, "table.pptx", stream.getvalue())
    assert search(store, "acceleration")[0].location == 1
    assert "Force | newtons" in search(store, "kilograms")[0].text
