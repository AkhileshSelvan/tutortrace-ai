"""Generate original educational fixtures, without external course material."""

from io import BytesIO
from pathlib import Path

import pymupdf
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
PDF_PAGES = (
    "Energy fundamentals\nKinetic energy is the energy of motion. The formula is K = 0.5 * mass * speed squared.\nFor mass 2 kg and speed 3 m/s, kinetic energy is 9 joules.",
    "Gravitational potential energy\nNear Earth's surface, gravitational potential energy is U = mass * g * height.\nUse g = 9.8 m/s squared. A 1 kg book raised by 2 m gains 19.6 joules.",
)
PPTX_SLIDES = (
    "Conservation of energy\nIn an isolated system, total energy remains constant. Energy can transfer between kinetic and potential forms.",
    "Power and energy transfer\nPower is energy transferred per unit time: power = energy / time.\nA lamp transferring 120 joules in 6 seconds has a power of 20 watts.",
)


def make_pdf(pages=PDF_PAGES) -> bytes:
    with pymupdf.open() as document:
        for text in pages:
            page = document.new_page(width=612, height=792)
            if text:
                page.insert_textbox(pymupdf.Rect(48, 48, 564, 744), text, fontsize=13)
        document.set_metadata({"title": "TutorTrace sample: Energy", "author": "TutorTrace AI demo fixtures"})
        return document.tobytes(no_new_id=True)


def make_pptx(slides=PPTX_SLIDES) -> bytes:
    presentation = Presentation()
    presentation.core_properties.title = "TutorTrace sample: Energy"
    presentation.core_properties.author = "TutorTrace AI demo fixtures"
    for text in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        if text:
            slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(5)).text_frame.text = text
    stream = BytesIO()
    presentation.save(stream)
    return stream.getvalue()


def main():
    destination = ROOT / "demo_assets"
    destination.mkdir(exist_ok=True)
    for name, data in (("energy_notes.pdf", make_pdf()), ("energy_slides.pptx", make_pptx())):
        path = destination / name
        path.write_bytes(data)
        print(f"Created {path.relative_to(ROOT)} ({len(data)} bytes)")


if __name__ == "__main__":
    main()
