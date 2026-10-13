"""Extract text boxes, grouped shapes and tables; never renumber blank slides."""

from io import BytesIO
from zipfile import ZipFile

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from tutortrace.errors import DocumentLimitError, InvalidDocumentError
from tutortrace.ingestion.pdf import MAX_TEXT_CHARS, MAX_UNITS
from tutortrace.models import ExtractedUnit, Extraction

MAX_ARCHIVE_ENTRIES = 10_000
MAX_EXPANDED_BYTES = 200 * 1024 * 1024


def _shape_text(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _shape_text(shape.shapes)
        elif shape.has_table:
            for row in shape.table.rows:
                yield " | ".join(cell.text for cell in row.cells if not cell.is_spanned)
        elif shape.has_text_frame:
            yield shape.text_frame.text


def extract_pptx(data: bytes) -> Extraction:
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ARCHIVE_ENTRIES or sum(info.file_size for info in entries) > MAX_EXPANDED_BYTES:
                raise DocumentLimitError("PowerPoint archive exceeds safe extraction limits (10,000 entries / 200 MiB expanded).")
            if "ppt/presentation.xml" not in archive.namelist():
                raise InvalidDocumentError("This ZIP file is not a PowerPoint presentation.")
        presentation = Presentation(BytesIO(data))
        if len(presentation.slides) > MAX_UNITS:
            raise DocumentLimitError(f"Presentation exceeds the {MAX_UNITS}-slide limit.")
        units = []
        total = 0
        for number, slide in enumerate(presentation.slides, 1):
            text = " ".join("\n".join(_shape_text(slide.shapes)).split())
            total += len(text)
            if total > MAX_TEXT_CHARS:
                raise DocumentLimitError("Extracted text exceeds the 5-million-character limit.")
            units.append(ExtractedUnit(number, text))
        warnings = ["Text boxes and tables only: images, diagrams, charts, speaker notes and animations are not interpreted."]
        blank = sum(not unit.text for unit in units)
        if blank:
            warnings.append(f"{blank} slide(s) have no extractable text; their slide numbers are retained.")
        return Extraction(tuple(units), tuple(warnings))
    except (InvalidDocumentError, DocumentLimitError):
        raise
    except Exception as exc:
        raise InvalidDocumentError("PowerPoint could not be read. Upload a valid .pptx file (legacy .ppt is unsupported).") from exc
